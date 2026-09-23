from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\Nevo\Downloads\financial material\Publishing and Video Production\Leadership and Governance")
SOURCE = ROOT / "Key-Person Risk in Sole Proprietorships and Founder-Led Businesses.md"
OUTPUT = ROOT / "Key-Person Risk in Sole Proprietorships and Founder-Led Businesses - IEEE Edition.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_cell_border(cell, color='D9D9D9'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = qn(f'w:{edge}')
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), color)


def plain_markdown(text):
    def link(match):
        return f"{match.group(1)} ({match.group(2)})"
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', link, text)
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'`([^`]+)`', r'\1', text)
    return text.strip()


def add_paragraph(doc, text, style=None, bold_lead=False):
    para = doc.add_paragraph(style=style)
    para.paragraph_format.space_after = Pt(7)
    para.paragraph_format.line_spacing = 1.15
    if bold_lead and ':' in text:
        lead, rest = text.split(':', 1)
        run = para.add_run(lead + ':')
        run.bold = True
        para.add_run(rest)
    else:
        para.add_run(plain_markdown(text))
    return para


def add_reference(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.25)
    para.paragraph_format.first_line_indent = Inches(-0.25)
    para.paragraph_format.space_after = Pt(5)
    para.paragraph_format.line_spacing = 1.1
    para.add_run(plain_markdown(text))
    return para


def add_table(doc, rows):
    headers = [item.strip() for item in rows[0].strip('|').split('|')]
    data = [[item.strip() for item in row.strip('|').split('|')] for row in rows[2:]]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header_cells = table.rows[0].cells
    for idx, label in enumerate(headers):
        header_cells[idx].text = plain_markdown(label)
        header_cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(header_cells[idx], '17365D')
        set_cell_margins(header_cells[idx])
        set_cell_border(header_cells[idx])
        for run in header_cells[idx].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            run.font.size = Pt(9)
    for row_index, values in enumerate(data):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = plain_markdown(value)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[idx])
            set_cell_border(cells[idx])
            if row_index % 2:
                set_cell_shading(cells[idx], 'F3F6FA')
            for paragraph in cells[idx].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size in [('Title', 20), ('Heading 1', 14), ('Heading 2', 12)]:
        style = doc.styles[style_name]
        style.font.name = 'Aptos Display' if style_name == 'Title' else 'Aptos'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(16 if style_name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(7)

    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    index = 0
    paragraph_buffer = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            text = ' '.join(line.strip() for line in paragraph_buffer).strip()
            if text:
                add_paragraph(doc, text)
            paragraph_buffer = []

    while index < len(lines):
        line = lines[index]
        if line.startswith('|'):
            flush_paragraph()
            table_rows = []
            while index < len(lines) and lines[index].startswith('|'):
                table_rows.append(lines[index])
                index += 1
            add_table(doc, table_rows)
            continue
        if not line.strip():
            flush_paragraph()
            index += 1
            continue
        if line.startswith('# '):
            flush_paragraph()
            para = doc.add_paragraph(style='Title')
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.add_run(plain_markdown(line[2:]))
        elif line.startswith('## '):
            flush_paragraph()
            doc.add_paragraph(plain_markdown(line[3:]), style='Heading 1')
        elif re.match(r'^\[\d+\]\s+', line):
            flush_paragraph()
            add_reference(doc, line)
        elif line.startswith('- '):
            flush_paragraph()
            para = doc.add_paragraph(style='List Bullet')
            para.paragraph_format.space_after = Pt(3)
            para.add_run(plain_markdown(line[2:]))
        else:
            paragraph_buffer.append(line)
        index += 1

    flush_paragraph()
    core = doc.core_properties
    core.title = 'Key Person Risk in Sole Proprietorships and Founder Led Businesses'
    core.subject = 'Leadership and governance article'
    core.author = 'SokoIntel'
    core.comments = 'Standalone leadership article'
    doc.save(OUTPUT)
    print(OUTPUT)

if __name__ == '__main__':
    main()