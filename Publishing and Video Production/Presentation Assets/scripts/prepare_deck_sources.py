"""Build portable slide-source metadata and PNG crops from saved Excel PDF exports."""
from pathlib import Path
import json, re, csv, hashlib
import pypdfium2 as pdfium
from pypdf import PdfReader
from PIL import Image, ImageChops
import openpyxl

ROOT=Path(__file__).resolve().parents[3]
ASSETS=ROOT/'Publishing and Video Production'/'Presentation Assets'
CATALOG=ROOT.parent/'business and technical website blog'/'content'/'catalog.json'
OUT=ASSETS/'content'
OUT.mkdir(parents=True,exist_ok=True)

def clean(s):
    s=re.sub(r'\[\[?(\d+)\]?\]\([^)]*\)',r'[\1]',s)
    s=re.sub(r'\[([^\]]+)\]\([^)]*\)',r'\1',s)
    return s.replace('**','').replace('*','').replace('`','').strip()

cat=json.loads(CATALOG.read_text(encoding='utf-8-sig'))
capture=json.loads((ASSETS/'Evidence/Workbooks/capture-manifest.json').read_text(encoding='utf-8-sig'))
workbooks={}
for cap in capture:
    aid=cap['asset_id']; pdf=ASSETS/'Evidence/Workbooks'/f'{aid}.pdf'
    doc=pdfium.PdfDocument(str(pdf)); pages=[]
    for i in range(len(doc)):
        img=doc[i].render(scale=2.1).to_pil().convert('RGB')
        bg=Image.new('RGB',img.size,'white')
        box=ImageChops.difference(img,bg).convert('L').point(lambda x:255 if x>22 else 0).getbbox()
        if box: img=img.crop((max(0,box[0]-12),max(0,box[1]-12),min(img.width,box[2]+12),min(img.height,box[3]+12)))
        dest=ASSETS/'Evidence/Workbooks'/f'{aid}-p{i+1}.png'; img.save(dest,optimize=True)
        pages.append(str(dest.relative_to(ASSETS)).replace('\\','/'))
    cap['images']=pages
    cap['pdf_text']='\n'.join(p.extract_text() or '' for p in PdfReader(str(pdf)).pages)
    path=cap['workbook']
    if path not in workbooks: workbooks[path]=openpyxl.load_workbook(path,read_only=True,data_only=True)
    ws=workbooks[path][cap['sheet']]
    cap['cells']={c.coordinate:str(c.value) for row in ws[cap['range']] for c in row if c.value is not None}
    cap['pdf']=str(pdf.relative_to(ASSETS)).replace('\\','/')

rows=list(csv.DictReader((ROOT/'Publishing and Video Production/Slide and Research Question Index.csv').open(encoding='utf-8-sig')))
articles=[]
for a in cat['articles']:
    body=a['body']; chunks=re.split(r'(?m)^#{2,3}\s+',body)
    sections=[]
    for chunk in chunks:
        heading,_,text=chunk.partition('\n')
        if len(sections)==0 and not body.startswith('#'): heading='Opening'; text=chunk
        sections.append({'heading':clean(heading),'text':clean(text)})
    refs=[s for s in sections if re.search(r'^References',s['heading'],re.I)]
    links=[]
    for label,url in re.findall(r'\[([^\]]+)\]\((https?://[^\s)]+)\)',body):
        if url not in [x['url'] for x in links]: links.append({'title':clean(label),'url':url})
    articles.append({k:a.get(k) for k in ['id','series_id','slug','title','short_title','questions','workbook_name']}|{'sections':sections,'references':refs,'links':links,'storyboard':[r for r in rows if r['Episode']==a['id']],'evidence':[c for c in capture if c['episode']==a['id']]})

payload={'source_catalog':str(CATALOG),'source_catalog_sha256':hashlib.sha256(CATALOG.read_bytes()).hexdigest(),'series':cat['series'],'research_date':cat.get('research_date'),'articles':articles}
(OUT/'source-material.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'source-digest.txt').write_text('\n\n'.join(a['id']+' '+a['title']+'\n'+'\n'.join(s['heading']+'\n'+s['text'][:1500] for s in a['sections'] if re.search(r'work|calculation|example|case|compare|fee|cost|follow|contribut|cash|ratio|tax|dividend|return|yield',s['heading'],re.I)) for a in articles),encoding='utf-8')
print(json.dumps({'articles':len(articles),'workbook_ranges':len(capture),'pngs':sum(len(c['images']) for c in capture),'output':str(OUT/'source-material.json')}))
