"""Verify punctuation cleanup, source reconstruction and Word/PDF output integrity."""
from pathlib import Path
from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
MASTER = ROOT / "NSE_Dividend_Growth_DCA_and_Watchlist_IEEE_source.md"
PART_1 = ROOT / "NSE_Dividend_Growth_Part_1_Core_Portfolio_and_DCA.md"
PART_2 = ROOT / "NSE_Dividend_Growth_Part_2_Watchlist.md"
OUTPUT = ROOT / "Publication" / "Output"


def main():
    master = MASTER.read_text(encoding="utf-8").splitlines()
    one = PART_1.read_text(encoding="utf-8").splitlines()
    two = PART_2.read_text(encoding="utf-8").splitlines()
    if master != one + two[5:]:
        raise AssertionError("The two Markdown parts do not reconstitute the normalized master")
    for path in [MASTER, PART_1, PART_2]:
        if "—" in path.read_text(encoding="utf-8"):
            raise AssertionError(f"Em dash found in {path.name}")

    report = []
    for folder in sorted(path for path in OUTPUT.iterdir() if path.is_dir()):
        docx = next(folder.glob("*.docx"))
        pdf = next(folder.glob("*.pdf"))
        document = Document(docx)
        reader = PdfReader(pdf)
        if not document.paragraphs or not reader.pages:
            raise AssertionError(f"Empty publication: {folder.name}")
        if not all((page.extract_text() or "").strip() for page in reader.pages):
            raise AssertionError(f"Blank PDF page: {folder.name}")
        report.append(f"{folder.name}: {len(document.paragraphs)} Word paragraphs; {len(reader.pages)} PDF pages")
    print("\n".join(report))


if __name__ == "__main__":
    main()
